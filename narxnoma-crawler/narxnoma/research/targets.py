"""Read the product list the user wants researched.

Accepted shapes, picked by file extension:

    .xlsx / .xlsm   first sheet, header row; columns found by name (see ALIASES)
    .csv            same, UTF-8 (with or without BOM), delimiter sniffed
    .txt            one product name per line; "name;price" or "name<TAB>price" carries a price

Only the name is required. A price is "our price" -- the baseline every delta is measured
against; with none, the workbook still compares the markets with each other. A category is a
narxnoma slug (see `CATEGORY_ALIASES` for the words accepted); with none, the run-wide
`default_category` applies.
"""

import csv
import re
from pathlib import Path

from narxnoma.matching.keys import CATEGORIES
from narxnoma.research.model import Target

ALIASES = {
    "name": (
        "name",
        "product",
        "product name",
        "title",
        "item",
        "товар",
        "наименование",
        "название",
        "mahsulot",
        "nomi",
    ),
    "price": ("our price", "price", "our price (sum)", "наша цена", "цена", "narx", "baho"),
    "category": ("category", "категория", "turkum", "kategoriya"),
    "currency": ("currency", "валюта", "valyuta"),
    "note": ("note", "notes", "comment", "примечание", "комментарий"),
}

CATEGORY_ALIASES = {
    "smartphones": (
        "smartphones",
        "smartphone",
        "phone",
        "phones",
        "смартфон",
        "смартфоны",
        "телефон",
        "телефоны",
        "смартфонлар",
        "smartfon",
        "telefon",
    ),
    "televisions": (
        "televisions",
        "television",
        "tv",
        "tvs",
        "тв",
        "телевизор",
        "телевизоры",
        "televizor",
    ),
    "laptops": ("laptops", "laptop", "ноутбук", "ноутбуки", "noutbuk"),
    "acs": (
        "acs",
        "ac",
        "air conditioner",
        "air conditioners",
        "кондиционер",
        "кондиционеры",
        "konditsioner",
    ),
    "fridges": ("fridges", "fridge", "refrigerator", "холодильник", "холодильники", "muzlatgich"),
    "washers": (
        "washers",
        "washer",
        "washing machine",
        "стиральная машина",
        "стиральные машины",
        "kir yuvish mashinasi",
    ),
    "ovens": ("ovens", "oven", "духовой шкаф", "духовка", "духовые шкафы", "duxovka"),
}
_CAT_LOOKUP = {a: slug for slug, names in CATEGORY_ALIASES.items() for a in names}


def category_slug(value) -> str | None:
    """A narxnoma slug for whatever the user wrote, or None."""
    if value is None:
        return None
    v = str(value).strip().lower()
    if v in CATEGORIES:
        return v
    return _CAT_LOOKUP.get(v)


def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _map_columns(headers: list) -> dict[str, int]:
    found: dict[str, int] = {}
    normed = [_norm_header(h) for h in headers]
    for key, names in ALIASES.items():
        for want in names:
            if want in normed and key not in found:
                found[key] = normed.index(want)
    return found


def _to_price(v) -> float | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v > 0 else None
    s = re.sub(r"[^\d.,]", "", str(v).replace("\xa0", "").replace(" ", ""))
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if len(s.rsplit(",", 1)[1]) <= 2 else s.replace(",", "")
    try:
        p = float(s)
    except ValueError:
        return None
    return p if p > 0 else None


def _cell(row, cols, key):
    i = cols.get(key)
    return row[i] if i is not None and i < len(row) else None


def _rows_to_targets(headers: list, rows, default_category: str | None) -> list[Target]:
    cols = _map_columns(headers)
    if "name" not in cols:
        raise ValueError(
            f"no product-name column found in {list(headers)!r}; "
            f"name one of: {', '.join(ALIASES['name'])}"
        )
    known = set(cols.values())
    out, seen = [], set()
    for row in rows:
        name = str(_cell(row, cols, "name") or "").strip()
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        cur = str(_cell(row, cols, "currency") or "UZS").strip().upper()
        out.append(
            Target(
                name=name,
                category=category_slug(_cell(row, cols, "category")) or default_category,
                our_price=_to_price(_cell(row, cols, "price")),
                currency="USD" if cur in ("USD", "$", "ДОЛЛАР") else "UZS",
                note=str(_cell(row, cols, "note") or ""),
                extra={
                    str(headers[i]): row[i]
                    for i in range(min(len(headers), len(row)))
                    if i not in known and headers[i] and row[i] not in (None, "")
                },
            )
        )
    return out


def load_targets(path: str | Path, default_category: str | None = None) -> list[Target]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook  # optional dependency, only for xlsx input

        ws = load_workbook(p, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = list(next(it))
        return _rows_to_targets(headers, [list(r) for r in it], default_category)
    if ext == ".csv":
        text = p.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        rows = list(csv.reader(text.splitlines(), dialect))
        return _rows_to_targets(rows[0], rows[1:], default_category)
    if ext == ".txt":
        out, seen = [], set()
        for line in p.read_text(encoding="utf-8-sig").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = re.split(r"\t|;", line, maxsplit=1)
            name = parts[0].strip()
            if name.lower() in seen:
                continue
            seen.add(name.lower())
            out.append(
                Target(
                    name=name,
                    category=default_category,
                    our_price=_to_price(parts[1]) if len(parts) > 1 else None,
                )
            )
        return out
    raise ValueError(f"unsupported target file type {ext!r}; use .xlsx, .csv or .txt")
